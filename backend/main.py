from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

from .compare import compare_records
from .vlm import extract_document, ollama_health


ROOT = Path(__file__).resolve().parents[1]
STATIC = ROOT / "static"

app = FastAPI(title="ReForm Local VLM Comparator", version="0.3.0")
app.mount("/static", StaticFiles(directory=str(STATIC)), name="static")


class RecordsExportRequest(BaseModel):
    records: List[Dict[str, Any]]
    filename: str = "vlm_generated.xlsx"


class CompareRequest(BaseModel):
    source_records: List[Dict[str, Any]]
    master_records: List[Dict[str, Any]]
    index_field: str


class ComparisonExportRequest(BaseModel):
    summary: Dict[str, Any]
    rows: List[Dict[str, Any]]
    details: List[Dict[str, Any]] = []


def _safe_excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def _workbook_response(wb: Workbook, filename: str) -> StreamingResponse:
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _write_rows(ws, rows: List[Dict[str, Any]]) -> None:
    if not rows:
        ws.append(["NO_DATA"])
        return

    headers: List[str] = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                headers.append(key)

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([_safe_excel_value(row.get(h)) for h in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


@app.get("/api/health")
def health():
    return ollama_health()


@app.post("/api/extract")
async def extract(
    file: UploadFile = File(...),
    schema_hint: str = Form(""),
):
    try:
        raw = await file.read()
        return extract_document(file.filename or "document", raw, schema_hint)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/read-table")
async def read_table(file: UploadFile = File(...)):
    try:
        raw = await file.read()
        name = (file.filename or "").lower()

        if name.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(raw), dtype=object)
        elif name.endswith((".xlsx", ".xlsm", ".xltx")):
            df = pd.read_excel(io.BytesIO(raw), dtype=object)
        else:
            raise ValueError("File data harus CSV atau XLSX.")

        df = df.where(pd.notnull(df), None)
        records = df.to_dict(orient="records")
        return {
            "filename": file.filename,
            "columns": [str(c) for c in df.columns],
            "records": records,
        }
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/export-generated")
def export_generated(payload: RecordsExportRequest):
    wb = Workbook()
    ws = wb.active
    ws.title = "Generated Data"

    rows = [
        {k: v for k, v in row.items() if k != "_page"}
        for row in payload.records
    ]
    _write_rows(ws, rows)

    filename = payload.filename
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    return _workbook_response(wb, filename)


@app.post("/api/compare")
def compare(payload: CompareRequest):
    try:
        return compare_records(
            payload.source_records,
            payload.master_records,
            payload.index_field,
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/export-comparison")
def export_comparison(payload: ComparisonExportRequest):
    wb = Workbook()

    # Primary output: actual side-by-side field comparison.
    ws = wb.active
    ws.title = "Comparison"
    _write_rows(ws, payload.details)

    if payload.details and "status" in payload.details[0]:
        headers = [cell.value for cell in ws[1]]
        status_col = headers.index("status") + 1
        fills = {
            "SAME": PatternFill("solid", fgColor="C6EFCE"),
            "DIFFERENT": PatternFill("solid", fgColor="FFC7CE"),
            "NOT_FOUND": PatternFill("solid", fgColor="FFEB9C"),
            "INDEX_EMPTY": PatternFill("solid", fgColor="D9E1F2"),
        }
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=status_col)
            fill = fills.get(str(cell.value))
            if fill:
                cell.fill = fill

    record_ws = wb.create_sheet("Record Summary")
    _write_rows(record_ws, payload.rows)

    summary_ws = wb.create_sheet("Summary")
    summary_ws.append(["metric", "value"])
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)
    for key, value in payload.summary.items():
        summary_ws.append([key, value])

    return _workbook_response(wb, "reform_comparison_result.xlsx")
